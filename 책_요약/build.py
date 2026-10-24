import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# (분야, 제목, 저자(추정), 위치, 확신도, 비고)
rows = [
# ── 과학 ──
("과학","크리스퍼가 온다 (CRISPR: A Crack in Creation)","제니퍼 다우드나","1단(위) 좌","상",""),
("과학","코스모스 (Cosmos)","칼 세이건","1단 좌","상",""),
("과학","그림으로 읽는 뇌과학의 모든 것","박문호","1단 좌","상",""),
("과학","뇌, 욕망의 비밀을 풀다","한스-게오르크 호이젤","1단 좌","중",""),
("과학","특이점이 온다 (The Singularity Is Near)","레이 커즈와일","4단(맨아래) 우","중",""),
# ── 심리학 / 뇌 ──
("심리학","EQ 감성지능 (Emotional Intelligence)","대니얼 골먼","1단 좌","상",""),
("심리학","생각에 관한 생각 (Thinking, Fast and Slow)","대니얼 카너먼","1단 좌","상",""),
("심리학","생각의 속도 (빌 게이츠)","빌 게이츠","2단 우","중","제목 불확실"),
# ── 철학 / 인문 / 역사 ──
("철학·인문","사피엔스 (Sapiens)","유발 하라리","2단 좌","상",""),
("철학·인문","호모 데우스 (Homo Deus)","유발 하라리","2단 좌","상",""),
("철학·인문","넥서스 (Nexus)","유발 하라리","2단 좌","상",""),
("철학·인문","명상록","마르쿠스 아우렐리우스","2단 중","상",""),
("철학·인문","월든 (Walden)","헨리 데이비드 소로","2단 중","중",""),
("철학·인문","채근담","홍자성","2단 좌","중",""),
("철학·인문","정의란 무엇인가 (Justice)","마이클 샌델","2단 중","상","한/영 각 1권"),
("철학·인문","죽음이란 무엇인가 (Death)","셸리 케이건","2단 중","상","예일대 강의"),
("역사","세계사 편력 1·2·3","자와할랄 네루","4단 좌","상","3권 세트"),
("역사","삼국지 (三國志) 전권 세트","나관중","4단 좌","상","약 10권 세트"),
# ── 투자 / 돈 / 경제 ──
("투자·경제","원칙 (Principles)","레이 달리오","2단 좌","상",""),
("투자·경제","The Barefoot Investor","스콧 파페","1단 우","상","영문"),
("투자·경제","부자의 그릇","이즈미 마사토","1단 우","중",""),
("투자·경제","시크릿 (The Secret)","론다 번","1단 우","상",""),
("투자·경제","MONEY (머니)","토니 로빈스","1단 우","상","영문"),
("투자·경제","부자 아빠 가난한 아빠","로버트 기요사키","1단 우","중",""),
("투자·경제","제로 투 원 (Zero to One)","피터 틸","2단 우","상",""),
# ── 경영 / 마케팅 ──
("경영·마케팅","마케팅 설계자 (닷컴 시크릿)","러셀 브런슨","3단 좌","상",""),
("경영·마케팅","This Is Marketing","세스 고딘","3단 좌","상","영문"),
("경영·마케팅","포지셔닝 (Positioning)","알 리스·잭 트라우트","3단 좌","상",""),
("경영·마케팅","Tribes (트라이브즈)","세스 고딘","3단 좌","중",""),
("경영·마케팅","Sustainable Fashion","-","3단 좌","중","영문"),
("경영·마케팅","초격차","권오현","3단 좌","중",""),
("경영·마케팅","린 인 (Lean In)","셰릴 샌드버그","2단 우","상",""),
("경영·마케팅","비커밍 (Becoming)","미셸 오바마","3단 중","상",""),
# ── 자기계발 ──
("자기계발","Think Like (Aaron Sansoni)","애런 산소니","1단 좌","중","영문"),
("자기계발","퓨처 셀프 (Future Self)","벤저민 하디","3단 중","중",""),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "책 목록"

headers = ["번호","분야","제목 (원제)","저자(추정)","위치(선반)","판독 확신도","비고"]
ws.append(headers)

hfill = PatternFill("solid", fgColor="2F5496")
hfont = Font(color="FFFFFF", bold=True, size=11)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin,right=thin,top=thin,bottom=thin)

for c in ws[1]:
    c.fill=hfill; c.font=hfont; c.alignment=Alignment(horizontal="center",vertical="center"); c.border=border

conf_fill = {"상":PatternFill("solid",fgColor="E2EFDA"),
             "중":PatternFill("solid",fgColor="FFF2CC"),
             "하":PatternFill("solid",fgColor="FCE4D6")}

for i,r in enumerate(rows,1):
    ws.append([i]+list(r))
    row = ws[i+1]
    for c in row:
        c.border=border; c.alignment=Alignment(vertical="center",wrap_text=True)
    row[5].fill = conf_fill.get(r[4], PatternFill())
    row[5].alignment = Alignment(horizontal="center",vertical="center")
    row[0].alignment = Alignment(horizontal="center",vertical="center")

widths=[6,14,42,22,14,12,24]
for i,w in enumerate(widths,1):
    ws.column_dimensions[chr(64+i)].width=w

ws.freeze_panes="A2"
ws.auto_filter.ref=f"A1:G{len(rows)+1}"

wb.save("/sessions/vibrant-magical-dijkstra/mnt/outputs/책장_도서목록.xlsx")
print("saved", len(rows), "rows")
